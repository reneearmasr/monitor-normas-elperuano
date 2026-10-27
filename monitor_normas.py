"""
Monitor de Normas Legales - Diario Oficial El Peruano
=====================================================

Revisa https://diariooficial.elperuano.pe/Normas en un rango de fechas y
selecciona las normas que involucran a:

  1. ATU (Autoridad de Transporte Urbano para Lima y Callao) y transporte urbano
  2. Normatividad presupuestal (transferencias, créditos suplementarios, etc.)
  3. Inversiones (Invierte.pe, proyectos de inversión, APP, OxI, etc.)

Genera un resumen en Word (.docx) y, si está configurado en config.ini,
lo envía por Telegram (listado de normas + el Word adjunto).

Uso:
    python monitor_normas.py                    # normas de hoy
    python monitor_normas.py --dias 7           # últimos 7 días
    python monitor_normas.py --desde 01/10/2026 --hasta 22/10/2026
    python monitor_normas.py --no-enviar        # solo genera el Word
"""

import argparse
import configparser
import json
import os
import shutil
import re
import sys
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from html import escape
from pathlib import Path

try:
    # Usa los certificados de Windows (necesario si la red inspecciona HTTPS)
    import truststore
    truststore.inject_into_ssl()
except ImportError:
    pass

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Twips

BASE_DIR = Path(__file__).resolve().parent
HORA_LIMA = timezone(timedelta(hours=-5))  # Perú no tiene horario de verano
URL_PORTAL = "https://diariooficial.elperuano.pe/Normas"
URL_FILTRO = "https://diariooficial.elperuano.pe/Normas/Filtro"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/130.0 Safari/537.36",
    "Accept-Language": "es-PE,es;q=0.9",
}

DIR_WEB = BASE_DIR / "docs"  # panel publicado en GitHub Pages
CAT_ATU = "ATU y transporte urbano"


def normalizar(texto: str) -> str:
    """Minúsculas, sin tildes y con guiones como espacios."""
    texto = unicodedata.normalize("NFD", texto or "")
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"[\s\-–—]+", " ", texto.lower())


def patron_termino(termino: str) -> re.Pattern:
    """Convierte un término del panel en regex: 'presupuest*' = cualquier terminación."""
    t = normalizar(termino).strip()
    comodin = t.endswith("*")
    t = t.rstrip("*").strip()
    cuerpo = r"\s+".join(re.escape(p) for p in t.split())
    return re.compile(r"\b" + cuerpo + (r"\w*" if comodin else r"\b"))


# --------------------------------------------------------------------------
# Criterios de búsqueda: docs/config/palabras_clave.json (editable desde el panel)
# --------------------------------------------------------------------------
_criterios = json.loads((DIR_WEB / "config" / "palabras_clave.json").read_text(encoding="utf-8"))
CATEGORIAS = {cat: terms for cat, terms in _criterios["categorias"].items()}
CATEGORIAS.setdefault(CAT_ATU, [])
PATRONES = {cat: [patron_termino(t) for t in terms if t.strip()]
            for cat, terms in CATEGORIAS.items()}
# Actos administrativos de personal que no interesan, salvo que mencionen a la ATU
EXCLUIR = [normalizar(t).strip() for t in _criterios.get("excluir_si_sumilla_empieza_con", [])
           if t.strip()]
# Temas ajenos que se cuelan por nombres de empresas (p. ej. "Inversiones X S.A.C.")
EXCLUIR_CONTIENE = [patron_termino(t) for t in _criterios.get("excluir_si_sumilla_contiene", [])
                    if t.strip()]

# Menciones a la ATU dentro del texto completo (sobre texto original)
REGEX_ATU_TEXTO = re.compile(
    r"\bATU\b|Autoridad de Transporte Urbano( para Lima y Callao)?", re.IGNORECASE
)
REGEX_MONTO = re.compile(r"S/\s?\d{1,3}(?:[ .,'’]\d{3})*(?:[.,]\d{2})?(?:\s?(?:millones|mil))?")
# Inicio de la parte resolutiva ("SE RESUELVE:", "DECRETA:", "ACUERDA:", "Ha dado la Ley siguiente:"…)
REGEX_RESUELVE = re.compile(
    r"(?:SE\s+)?(?:RESUELVE|DECRETA|ACUERDA|ACUERDAN|DISPONE|RESOLVI[ÓO])\s*:"
    r"|ACORD[ÓO](?:\s+POR\s+\w+)?\s*:|Ha dado la Ley siguiente\s*:",
    re.IGNORECASE)
# Artículo 1 resolutivo: "Artículo 1.-", "Artículo Primero.-", "ARTÍCULO PRIMERO. -" (no las citas de los considerandos)
REGEX_ART1 = re.compile(r"Art[íi]culo\s+(?:1|Primero|[ÚU]nico)\s*[.º°o]*\s*[-–—:]", re.IGNORECASE)


@dataclass
class Norma:
    fecha: str
    entidad: str
    titulo: str
    sumilla: str
    url: str
    url_pdf: str
    extraordinaria: bool
    codigo: str = ""          # código de publicación, p. ej. 2568110-1
    digital: bool = False     # "Publicación Digital" (sobre todo municipalidades)
    texto: str = ""
    categorias: list = field(default_factory=list)
    menciones_atu: list = field(default_factory=list)


# --------------------------------------------------------------------------
# Descarga y lectura del portal
# --------------------------------------------------------------------------
def nueva_sesion() -> requests.Session:
    s = requests.Session()
    s.headers.update(HEADERS)
    s.get(URL_PORTAL, timeout=30)  # obtiene cookies del portal
    return s


def listar_normas(s: requests.Session, desde: date, hasta: date, solo_extra=False):
    params = {"dateparam": datetime.now().strftime("%m/%d/%Y 00:00:00")}
    data = {
        "cddesde": desde.strftime("%d/%m/%Y"),
        "cdhasta": hasta.strftime("%d/%m/%Y"),
        "btnBuscar": "Buscar",
    }
    if solo_extra:
        data["tipo"] = "on"
    for intento in range(4):
        try:
            r = s.post(URL_FILTRO, params=params, data=data, timeout=120,
                       headers={"X-Requested-With": "XMLHttpRequest"})
            r.raise_for_status()
            break
        except requests.RequestException as e:
            if intento == 3:
                raise
            print(f"  (reintento del listado: {e.__class__.__name__})", flush=True)
            time.sleep(10 * (intento + 1))
    soup = BeautifulSoup(r.text, "lxml")

    normas = []
    # Las "Publicación Digital" usan la clase edicionesoficiales_articulos_dig y enlaces de epdoc2
    for art in soup.select("article.edicionesoficiales_articulos, article.edicionesoficiales_articulos_dig"):
        cont = art.select_one(".ediciones_texto")
        if not cont:
            continue
        enlace = cont.select_one("h5 a")
        fecha_tag = cont.select_one("p b")
        parrafos = cont.find_all("p")
        pdf = art.select_one(".ediciones_botones a[href$='/pdf']")
        url = enlace["href"] if enlace else ""
        portada = art.select_one(".ediciones_pdf img")
        m = (re.search(r"/(\d+-\d+)(?:/|$)", url)
             or re.search(r"/(\d+-\d+)_Portada", portada["src"] if portada else ""))
        entidad = cont.h4.get_text(" ", strip=True) if cont.h4 else ""
        digital = "_dig" in " ".join(art.get("class", [])) or "Publicación Digital" in entidad
        normas.append(Norma(
            fecha=(fecha_tag.get_text(strip=True).replace("Fecha:", "").strip()
                   if fecha_tag else ""),
            entidad=re.sub(r"\s*Publicación Digital\s*$", "", entidad).strip(),
            titulo=enlace.get_text(" ", strip=True) if enlace else "",
            sumilla=parrafos[-1].get_text(" ", strip=True) if len(parrafos) > 1 else "",
            url=url,
            url_pdf=pdf["href"] if pdf else "",
            extraordinaria=bool(cont.select_one(".extraordinaria")),
            codigo=m.group(1) if m else "",
            digital=digital,
        ))
    return normas


def codigo_norma(norma: Norma) -> str:
    """Código de publicación, p. ej. 2567343-1, tomado del enlace del portal."""
    if norma.codigo:
        return norma.codigo
    m = re.search(r"/(\d+-\d+)(?:/|$)", norma.url)
    return m.group(1) if m else ""


def url_texto(norma: Norma) -> str:
    # Los enlaces /dispositivo/NL/... del portal suelen dar 404; este visor sí responde
    op = codigo_norma(norma)
    return f"https://busquedas.elperuano.pe/api/visor_html/{op}" if op else norma.url


def descargar_texto(s: requests.Session, norma: Norma, reintentos=2) -> str:
    op = codigo_norma(norma)
    if not op:
        return ""
    for intento in range(reintentos + 1):
        try:
            r = s.get(url_texto(norma), timeout=30)
            if r.status_code == 404:
                return ""
            r.raise_for_status()
            r.encoding = "utf-8"
            soup = BeautifulSoup(r.text, "lxml")
            nodo = soup.find(id=f"x{op}") or soup.select_one(".story") or soup.body
            return nodo.get_text("\n", strip=True) if nodo else ""
        except requests.RequestException:
            time.sleep(2 * (intento + 1))
    return ""


# --------------------------------------------------------------------------
# Clasificación y extractos
# --------------------------------------------------------------------------
def clasificar(norma: Norma):
    cabecera = normalizar(f"{norma.entidad} {norma.titulo} {norma.sumilla}")
    cats = [cat for cat, pats in PATRONES.items() if any(p.search(cabecera) for p in pats)]

    # La ATU puede aparecer solo en el texto o anexos (p. ej. transferencias del MEF);
    # en las normas que emite la propia ATU el extracto no aporta nada
    emitida_por_atu = "autoridad de transporte urbano" in normalizar(norma.entidad)
    if norma.texto and not emitida_por_atu:
        for m in REGEX_ATU_TEXTO.finditer(norma.texto):
            ini, fin = max(0, m.start() - 220), min(len(norma.texto), m.end() + 220)
            frag = re.sub(r"\s+", " ", norma.texto[ini:fin]).strip()
            if not any(frag[:80] in x for x in norma.menciones_atu):
                norma.menciones_atu.append(f"…{frag}…")
            if len(norma.menciones_atu) >= 3:
                break
    menciona_atu = bool(emitida_por_atu or norma.menciones_atu
                        or REGEX_ATU_TEXTO.search(norma.sumilla))
    sumilla = normalizar(norma.sumilla).strip()
    if not menciona_atu and (sumilla.startswith(tuple(EXCLUIR))
                             or any(p.search(sumilla) for p in EXCLUIR_CONTIENE)):
        cats = []
    if menciona_atu and CAT_ATU not in cats:
        cats.insert(0, CAT_ATU)
    norma.categorias = cats
    return bool(cats)


def extracto_articulos(texto: str, max_chars=700) -> str:
    """Primer artículo de la parte resolutiva (lo que la norma dispone)."""
    if not texto:
        return ""
    # Se busca el artículo 1 después de "SE RESUELVE:" para no tomar artículos citados en los considerandos
    marca = REGEX_RESUELVE.search(texto)
    inicio = marca.end() if marca else 0
    m = REGEX_ART1.search(texto, inicio) or (REGEX_ART1.search(texto) if inicio else None)
    if m:
        frag = texto[m.start():]
    elif inicio:
        frag = texto[inicio:]
    else:
        return ""
    frag = re.sub(r"\s+", " ", frag).strip()
    if len(frag) > max_chars:
        frag = frag[:max_chars].rsplit(" ", 1)[0] + "…"
    return frag


def montos(texto: str, n=4):
    """Montos en soles citados en la norma; descarta cifras menores a S/ 1 000 (numerales, anexos)."""
    vistos, claves = [], set()
    for m in REGEX_MONTO.findall(texto or ""):
        m = m.strip().rstrip(".,")
        cifra = re.sub(r"[.,]\d{2}$", "", re.sub(r"[^\d.,]", "", m))
        valor = int(re.sub(r"\D", "", cifra) or 0)
        # el mismo monto suele repetirse con otro formato (419,268.00 / 419 268,00)
        if (valor >= 1000 or "millon" in m) and valor not in claves:
            claves.add(valor)
            vistos.append(m)
        if len(vistos) >= n:
            break
    return vistos


# --------------------------------------------------------------------------
# Documento Word (informe formal, con la paleta del Portal de Gestión)
# --------------------------------------------------------------------------
NAVY, AZUL, GRIS, FONDO, BORDE, ROJO = "0A2F5C", "0B5CAD", "5B6880", "F2F5FA", "C9D2DF", "B42318"
COLOR_TEMA = {CAT_ATU: "2A78D6", "Normatividad presupuestal": "EB6834", "Inversiones": "1BAF7A"}
FUENTE = "Arial"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


def fecha_larga(d: date, con_dia=True) -> str:
    base = f"{d.day} de {MESES[d.month - 1]} de {d.year}"
    return f"{DIAS[d.weekday()]} {base}" if con_dia else base


def periodo_largo(desde: date, hasta: date) -> str:
    if desde == hasta:
        return fecha_larga(desde)
    if (desde.year, desde.month) == (hasta.year, hasta.month):
        return f"del {desde.day} al {fecha_larga(hasta, False)}"
    return f"del {fecha_larga(desde, False)} al {fecha_larga(hasta, False)}"


def _rgb(hexa):
    return RGBColor.from_string(hexa)


def _run(par, texto, size=10, bold=False, color=None, italic=False):
    r = par.add_run(texto)
    r.font.name = FUENTE
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FUENTE)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = _rgb(color)
    return r


def _par(cont, texto="", size=10, bold=False, color=None, align=None, antes=0, despues=4,
         italic=False, interlineado=1.15):
    p = cont.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(antes), Pt(despues), interlineado
    if align is not None:
        p.alignment = align
    if texto:
        _run(p, texto, size, bold, color, italic)
    return p


def _celda(celda, texto="", size=9, bold=False, color=None, align=None, italic=False):
    """Escribe en el primer párrafo de la celda (sin el espaciado del estilo Normal)."""
    p = celda.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(1), Pt(1), 1.1
    if align is not None:
        p.alignment = align
    if texto:
        _run(p, texto, size, bold, color, italic)
    return p


def _borde_parrafo(par, lado="bottom", color=NAVY, grosor=12, espacio=4):
    ppr = par._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    b = OxmlElement(f"w:{lado}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(grosor))
    b.set(qn("w:space"), str(espacio)); b.set(qn("w:color"), color)
    bdr.append(b)


def _bordes_tabla(tabla, color=BORDE, grosor=4, internos=True):
    tblpr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    lados = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if internos else [])
    for lado in lados:
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(grosor))
        b.set(qn("w:space"), "0"); b.set(qn("w:color"), color)
        bordes.append(b)
    tblpr.append(bordes)
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("top", 50), ("bottom", 50), ("left", 100), ("right", 100)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def _borde_celda(celda, lado, color, grosor):
    tcpr = celda._tc.get_or_add_tcPr()
    bordes = tcpr.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        tcpr.append(bordes)
    b = OxmlElement(f"w:{lado}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(grosor))
    b.set(qn("w:space"), "0"); b.set(qn("w:color"), color)
    bordes.append(b)


def sombrear(celda, color_hex):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color_hex)
    tcpr.append(shd)


def _anchos(tabla, anchos_cm):
    tabla.autofit = False
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_cm):
            celda.width = Cm(ancho)


def _fila_fija(fila, encabezado=False):
    trpr = fila._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if encabezado:
        trpr.append(OxmlElement("w:tblHeader"))


def _campo(par, instruccion, size=8, color=GRIS):
    """Campo de Word (PAGE / NUMPAGES) para la numeración de páginas."""
    r = _run(par, "", size, color=color)
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r._r.append(fc)
        elif texto == instruccion:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {texto} "; r._r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = texto; r._r.append(t)


def agregar_hipervinculo(parrafo, url, texto, size=9, bold=False, color=AZUL):
    rel = parrafo.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fuentes = OxmlElement("w:rFonts")
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fuentes.set(qn(k), FUENTE)
    rpr.append(fuentes)
    if bold:
        rpr.append(OxmlElement("w:b"))
    c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = texto; t.set(qn("xml:space"), "preserve"); run.append(t)
    hl.append(run)
    parrafo._p.append(hl)


def _solo_tab_derecha(par, ancho_cm):
    """Quita las tabulaciones del estilo Encabezado/Pie (centro 8,25 cm y derecha 16,5 cm) y deja una a la derecha."""
    ts = par.paragraph_format.tab_stops
    for pos in (Twips(4680), Twips(9360)):
        ts.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Cm(ancho_cm), WD_TAB_ALIGNMENT.RIGHT)


def _titulo_seccion(doc, texto):
    p = _par(doc, texto.upper(), 11, True, NAVY, antes=14, despues=6)
    _borde_parrafo(p, "bottom", BORDE, 6, 2)
    p.paragraph_format.keep_with_next = True
    return p


def nombre_propio(texto: str) -> str:
    """'AUTORIDAD DE TRANSPORTE URBANO PARA LIMA Y CALLAO' -> 'Autoridad de Transporte Urbano para Lima y Callao'."""
    menores = {"de", "del", "la", "las", "los", "el", "y", "e", "para", "en", "a", "con", "por"}
    palabras = texto.lower().split()
    return " ".join(w if i and w in menores else w[:1].upper() + w[1:] for i, w in enumerate(palabras))


def _plural(n, sing, plur):
    return f"{n} {sing if n == 1 else plur}"


def generar_word(seleccion, total, desde, hasta, ruta: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2)
    sec.header_distance = sec.footer_distance = Cm(1)
    ancho = 21 - 4.4
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FUENTE, Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FUENTE)

    # Encabezado y pie de página
    cab = sec.header.paragraphs[0]
    _solo_tab_derecha(cab, ancho)
    _run(cab, "AUTORIDAD DE TRANSPORTE URBANO PARA LIMA Y CALLAO", 7.5, True, NAVY)
    _run(cab, "\tMonitor de Normas Legales", 7.5, color=GRIS)
    _borde_parrafo(cab, "bottom", NAVY, 8, 4)
    pie = sec.footer.paragraphs[0]
    _solo_tab_derecha(pie, ancho)
    _borde_parrafo(pie, "top", BORDE, 6, 4)
    _run(pie, "Fuente: Diario Oficial El Peruano · Elaboración automática", 7.5, color=GRIS)
    _run(pie, "\tPágina ", 7.5, color=GRIS)
    _campo(pie, "PAGE", 7.5)
    _run(pie, " de ", 7.5, color=GRIS)
    _campo(pie, "NUMPAGES", 7.5)

    # Título
    _par(doc, "RESUMEN DE NORMAS LEGALES", 18, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 0, 2)
    _par(doc, f"Diario Oficial El Peruano · {periodo_largo(desde, hasta).capitalize()}",
         10.5, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, despues=2)
    regla = _par(doc, "", 4, despues=10)
    _borde_parrafo(regla, "bottom", NAVY, 18, 1)

    # Ficha técnica
    por_tema = {c: [n for n in seleccion if n.categorias and n.categorias[0] == c] for c in CATEGORIAS}
    ficha = doc.add_table(rows=2, cols=4)
    _bordes_tabla(ficha)
    datos = [("Periodo revisado", desde.strftime("%d/%m/%Y") if desde == hasta
              else f"{desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"),
             ("Fecha de elaboración", datetime.now(HORA_LIMA).strftime("%d/%m/%Y %H:%M")),
             ("Normas publicadas", str(total)),
             ("Normas de interés", str(len(seleccion)))]
    for i, (lbl, val) in enumerate(datos):
        fila, col = ficha.rows[i // 2], (i % 2) * 2
        sombrear(fila.cells[col], FONDO)
        _celda(fila.cells[col], lbl, 8.5, True, NAVY)
        _celda(fila.cells[col + 1], val, 9.5, i == 3, NAVY if i == 3 else None)
    _anchos(ficha, [3.6, 4.7, 3.6, 4.7])

    # 1. Resumen ejecutivo
    _titulo_seccion(doc, "1. Resumen ejecutivo")
    if seleccion:
        etiqueta = {CAT_ATU: "sobre la ATU y el transporte urbano"}
        partes = [f"{len(por_tema[c])} {etiqueta.get(c, 'de ' + c[0].lower() + c[1:])}"
                  for c in CATEGORIAS if por_tema[c]]
        lista = partes[0] if len(partes) == 1 else f"{', '.join(partes[:-1])} y {partes[-1]}"
        texto = (f"De las {total} normas publicadas en el periodo, "
                 f"{_plural(len(seleccion), 'resulta', 'resultan')} de interés institucional: {lista}. "
                 + ("" if por_tema[CAT_ATU] else "No se identificaron normas que involucren a la ATU. ")
                 + "El Cuadro 1 presenta la relación completa y la sección 2, el detalle de cada norma.")
    else:
        texto = (f"De las {total} normas publicadas en el periodo, ninguna corresponde a la ATU, "
                 "a la normatividad presupuestal ni a inversiones.")
    _par(doc, texto, 10, despues=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    atu = [n for n in seleccion if CAT_ATU in n.categorias]
    if atu:
        caja = doc.add_table(rows=1, cols=1)
        _bordes_tabla(caja, "F4C7C3", 4, internos=False)
        c = caja.rows[0].cells[0]
        sombrear(c, "FDF3F2")
        _borde_celda(c, "left", ROJO, 24)
        p = _celda(c, f"ATENCIÓN: {_plural(len(atu), 'norma involucra', 'normas involucran')} "
                      "directamente a la ATU", 9, True, ROJO)
        for n in atu:
            p = _par(c, "", despues=1, interlineado=1.1)
            _run(p, "▪ ", 9, color=ROJO)
            _run(p, f"{n.titulo}", 9, True)
            _run(p, f" ({nombre_propio(n.entidad)}): {n.sumilla}", 9)
        _anchos(caja, [ancho])
        _par(doc, "", 4, despues=4)

    if seleccion:
        cap = _par(doc, "Cuadro 1. Normas de interés identificadas", 9, True, NAVY, antes=4, despues=3)
        cap.paragraph_format.keep_with_next = True
        cuadro = doc.add_table(rows=1, cols=4)
        _bordes_tabla(cuadro)
        enc = cuadro.rows[0]
        _fila_fija(enc, encabezado=True)
        for celda, txt in zip(enc.cells, ["N.°", "Norma y entidad emisora", "Sumilla", "Publicación"]):
            sombrear(celda, NAVY)
            _celda(celda, txt, 8.5, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        num = 0
        for cat, normas_cat in por_tema.items():
            if not normas_cat:
                continue
            grupo = cuadro.add_row()
            _fila_fija(grupo)
            g = grupo.cells[0].merge(grupo.cells[3])
            sombrear(g, FONDO)
            _borde_celda(g, "left", COLOR_TEMA.get(cat, AZUL), 30)
            _celda(g, f"{cat.upper()}  ({len(normas_cat)})", 8.5, True, NAVY)
            for n in normas_cat:
                num += 1
                fila = cuadro.add_row()
                _fila_fija(fila)
                c0, c1, c2, c3 = fila.cells
                _celda(c0, str(num), 9, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER)
                agregar_hipervinculo(_celda(c1), url_texto(n), n.titulo, 8.5, True, AZUL)
                _par(c1, n.entidad, 7.5, color=GRIS, despues=1, interlineado=1.05)
                _celda(c2, n.sumilla, 8.5)
                _celda(c3, n.fecha, 8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                if n.digital and not n.extraordinaria:
                    _par(c3, "Publicación digital", 7, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER,
                         despues=0, interlineado=1.0)
                if n.extraordinaria:
                    _par(c3, "Edición extraordinaria", 7, color=ROJO, align=WD_ALIGN_PARAGRAPH.CENTER,
                         despues=0, interlineado=1.0)
        _anchos(cuadro, [1.0, 5.0, 8.4, 2.2])

        # 2. Detalle
        _titulo_seccion(doc, "2. Detalle de las normas")
        num, sub = 0, 0
        for cat, normas_cat in por_tema.items():
            if not normas_cat:
                continue
            sub += 1
            p = _par(doc, f"2.{sub}  {cat}", 10.5, True, AZUL, antes=8, despues=4)
            p.paragraph_format.keep_with_next = True
            for n in normas_cat:
                num += 1
                ficha = doc.add_table(rows=1, cols=2)
                _bordes_tabla(ficha)
                cab_f = ficha.rows[0]
                _fila_fija(cab_f)
                t = cab_f.cells[0].merge(cab_f.cells[1])
                sombrear(t, "E8F0FA")
                _borde_celda(t, "left", COLOR_TEMA.get(cat, AZUL), 30)
                _celda(t, f"{num}.  {n.titulo}", 9.5, True, NAVY)

                def fila_dato(etiqueta, valor=None, color_lbl=NAVY):
                    f = ficha.add_row()
                    _fila_fija(f)
                    sombrear(f.cells[0], FONDO)
                    _celda(f.cells[0], etiqueta, 8.5, True, color_lbl)
                    if valor is not None:
                        _celda(f.cells[1], valor, 9)
                    return f.cells[1]

                fila_dato("Entidad emisora", n.entidad)
                fila_dato("Publicación", f"{n.fecha} · "
                          + ("Edición extraordinaria" if n.extraordinaria
                             else "Publicación digital" if n.digital else "Edición ordinaria"))
                fila_dato("Temas", " · ".join(n.categorias))
                fila_dato("Sumilla", n.sumilla)
                disp = extracto_articulos(n.texto)
                if disp:
                    fila_dato("Disposición principal", disp)
                mts = montos(n.texto)
                if mts:
                    fila_dato("Montos citados", " · ".join(mts))
                if n.menciones_atu:
                    c = fila_dato("Mención a la ATU", color_lbl=ROJO)
                    for i, frag in enumerate(n.menciones_atu):
                        p = _celda(c) if i == 0 else _par(c, "", despues=1, interlineado=1.1)
                        _run(p, frag, 8.5, italic=True, color=GRIS)
                c = fila_dato("Enlaces")
                p = _celda(c)
                agregar_hipervinculo(p, url_texto(n), "Texto completo", 8.5)
                _run(p, "   ·   ", 8.5, color=GRIS)
                agregar_hipervinculo(p, n.url, "El Peruano", 8.5)
                if n.url_pdf:
                    _run(p, "   ·   ", 8.5, color=GRIS)
                    agregar_hipervinculo(p, n.url_pdf, "PDF", 8.5)
                _anchos(ficha, [3.6, ancho - 3.6])
                _par(doc, "", 4, despues=6)

    nota = _par(doc, "", antes=10, despues=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _borde_parrafo(nota, "top", BORDE, 4, 4)
    _run(nota, "Nota metodológica: ", 7.5, True, GRIS)
    _run(nota, "se revisan todas las normas publicadas en el Diario Oficial El Peruano en el periodo. "
               "Una norma se selecciona cuando su entidad, título o sumilla contiene alguno de los términos "
               "definidos para cada tema, o cuando su texto completo menciona a la ATU. Los criterios pueden "
               "consultarse y editarse en el panel: ", 7.5, color=GRIS)
    agregar_hipervinculo(nota, "https://reneearmasr.github.io/monitor-normas-elperuano/",
                         "reneearmasr.github.io/monitor-normas-elperuano", 7.5, color=AZUL)

    doc.save(ruta)


# --------------------------------------------------------------------------
# Envío
# --------------------------------------------------------------------------
def texto_breve(seleccion, total, desde, hasta):
    """Listado en HTML de Telegram, partido en mensajes de menos de 4096 caracteres."""
    periodo = (f"{desde:%d/%m/%Y}" if desde == hasta else f"{desde:%d/%m/%Y} al {hasta:%d/%m/%Y}")
    lineas = [f"<b>Normas El Peruano – {periodo}</b>",
              f"Revisadas: {total} | De interés: {len(seleccion)}"]
    for cat in CATEGORIAS:
        normas_cat = [n for n in seleccion if n.categorias and n.categorias[0] == cat]
        if normas_cat:
            lineas.append(f"\n<b>{escape(cat)} ({len(normas_cat)})</b>")
            for n in normas_cat:
                lineas.append(f'• <a href="{escape(url_texto(n))}">{escape(n.titulo)}</a> '
                              f"({escape(n.entidad)}): {escape(n.sumilla[:200])}")
    mensajes, actual = [], ""
    for linea in lineas:
        if len(actual) + len(linea) > 3800:
            mensajes.append(actual)
            actual = ""
        actual += linea + "\n"
    mensajes.append(actual)
    return mensajes


def enviar_telegram(cfg, ruta_docx: Path, mensajes, titulo: str):
    t = cfg["telegram"]
    token = os.environ.get("MONITOR_TELEGRAM_TOKEN") or t.get("bot_token", "")
    chat_id = os.environ.get("MONITOR_TELEGRAM_CHAT_ID") or t.get("chat_id", "")
    if not (token and chat_id):
        print("[telegram] Falta bot_token o chat_id en config.ini; no se envía.")
        return
    api = f"https://api.telegram.org/bot{token}"
    for texto in mensajes:
        r = requests.post(f"{api}/sendMessage", timeout=60, data={
            "chat_id": chat_id, "text": texto, "parse_mode": "HTML",
            "disable_web_page_preview": "true"})
        if not r.ok:
            print(f"[telegram] Error {r.status_code}: {r.text[:300]}")
            return
    with open(ruta_docx, "rb") as fh:
        r = requests.post(f"{api}/sendDocument", timeout=120,
                          data={"chat_id": chat_id, "caption": titulo},
                          files={"document": (ruta_docx.name, fh)})
    if not r.ok:
        print(f"[telegram] Error al enviar el Word {r.status_code}: {r.text[:300]}")
        return
    print(f"[telegram] Enviado: {len(mensajes)} mensaje/s + Word")


# --------------------------------------------------------------------------
# --------------------------------------------------------------------------
# Historial del panel web
# --------------------------------------------------------------------------
def fecha_iso(fecha: str) -> str:
    try:
        return datetime.strptime(fecha, "%d/%m/%Y").strftime("%Y-%m-%d")
    except ValueError:
        return ""


def publicar_historial(normas, seleccion, ruta_docx: Path, hasta: date):
    """Agrega la revisión a docs/data/normas.json y copia el Word a docs/resumenes."""
    destino = DIR_WEB / "resumenes" / ruta_docx.name
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ruta_docx, destino)

    archivo = DIR_WEB / "data" / "normas.json"
    archivo.parent.mkdir(parents=True, exist_ok=True)
    datos = (json.loads(archivo.read_text(encoding="utf-8")) if archivo.exists()
             else {"normas": [], "revisadas_por_dia": {}})

    # Copia con nombre fijo del resumen más reciente: la usa el envío por correo (Power Automate)
    if hasta.isoformat() >= datos.get("ultimo_resumen", {}).get("fecha", ""):
        shutil.copyfile(ruta_docx, DIR_WEB / "resumenes" / "Normas_ElPeruano_ultimo.docx")
        datos["ultimo_resumen"] = {"fecha": hasta.isoformat(), "archivo": f"resumenes/{ruta_docx.name}",
                                   "normas_de_interes": len(seleccion)}

    # Total de normas revisadas por fecha de publicación (para los indicadores)
    por_dia = {}
    for n in normas:
        f = fecha_iso(n.fecha)
        if f:
            por_dia[f] = por_dia.get(f, 0) + 1
    # max(): en modo --diario la fecha de ayer solo trae sus extraordinarias
    rpd = datos["revisadas_por_dia"]
    for f, c in por_dia.items():
        rpd[f] = max(rpd.get(f, 0), c)

    nuevas = {}
    for n in seleccion:
        nuevas[codigo_norma(n) or n.url] = {
            "id": codigo_norma(n) or n.url,
            "fecha": fecha_iso(n.fecha),
            "entidad": n.entidad,
            "titulo": n.titulo,
            "sumilla": n.sumilla,
            "categorias": n.categorias,
            "extraordinaria": n.extraordinaria,
            "digital": n.digital,
            "dispone": extracto_articulos(n.texto),
            "montos": montos(n.texto),
            "menciones_atu": n.menciones_atu,
            "url": n.url,
            "url_texto": url_texto(n),
            "url_pdf": n.url_pdf,
            "resumen": f"resumenes/{ruta_docx.name}",
        }
    conservar = [x for x in datos["normas"] if x["id"] not in nuevas]
    datos["normas"] = sorted(conservar + list(nuevas.values()),
                             key=lambda x: (x["fecha"], x["entidad"]), reverse=True)
    datos["actualizado"] = datetime.now(HORA_LIMA).strftime("%Y-%m-%dT%H:%M:%S-05:00")
    datos["categorias"] = list(CATEGORIAS)
    archivo.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"  Historial del panel: {len(datos['normas'])} normas ({archivo})")


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Monitor de normas El Peruano para la ATU")
    ap.add_argument("--desde", help="dd/mm/aaaa (por defecto: hoy)")
    ap.add_argument("--hasta", help="dd/mm/aaaa (por defecto: hoy)")
    ap.add_argument("--dias", type=int, help="revisar los últimos N días (incluye hoy)")
    ap.add_argument("--sin-texto", action="store_true",
                    help="no descargar el texto completo (más rápido, pero no detecta "
                         "menciones a la ATU dentro de anexos)")
    ap.add_argument("--no-enviar", action="store_true", help="solo genera el Word")
    ap.add_argument("--chat-id", action="store_true",
                    help="muestra el chat_id de quien escribió al bot y termina")
    ap.add_argument("--publicar", action="store_true",
                    help="agrega las normas al historial del panel web (docs/data)")
    ap.add_argument("--diario", action="store_true",
                    help="normas de hoy más las ediciones extraordinarias de ayer "
                         "(publicadas de noche, después de la revisión del día anterior)")
    args = ap.parse_args()

    hoy = datetime.now(HORA_LIMA).date()
    hasta = datetime.strptime(args.hasta, "%d/%m/%Y").date() if args.hasta else hoy
    if args.dias:
        desde = hasta - timedelta(days=args.dias - 1)
    else:
        desde = datetime.strptime(args.desde, "%d/%m/%Y").date() if args.desde else hasta

    cfg = configparser.ConfigParser()
    cfg.read(BASE_DIR / "config.ini", encoding="utf-8")
    for sec in ("telegram", "general"):
        if not cfg.has_section(sec):
            cfg.add_section(sec)

    if args.chat_id:
        token = os.environ.get("MONITOR_TELEGRAM_TOKEN") or cfg["telegram"].get("bot_token", "")
        if not token:
            print("Primero coloque bot_token en config.ini.")
            return 1
        api = f"https://api.telegram.org/bot{token}"
        yo = requests.get(f"{api}/getMe", timeout=30).json()
        info = yo.get("result", {})
        print(f"Bot: @{info.get('username')} ({info.get('first_name')})  {yo.get('description', '')}")
        r = requests.get(f"{api}/getUpdates", timeout=30).json()
        if not r.get("ok"):
            print(f"Telegram respondió: {r.get('description')}")
        chats = {}
        for u in r.get("result", []):
            chat = (u.get("message") or u.get("channel_post") or {}).get("chat")
            if chat:
                chats[chat["id"]] = chat.get("title") or chat.get("first_name", "")
        if not chats:
            print("No hay mensajes: envíe un mensaje a su bot y vuelva a ejecutar.")
        for cid, nombre in chats.items():
            print(f"chat_id = {cid}   ({nombre})")
        return 0

    print(f"Consultando normas del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}…")
    s = nueva_sesion()
    normas = listar_normas(s, desde, hasta)
    if args.diario:
        urls = {n.url for n in normas}
        ayer = desde - timedelta(days=1)
        extra = [n for n in listar_normas(s, ayer, ayer, solo_extra=True)
                 if n.extraordinaria and n.url not in urls]
        print(f"  + {len(extra)} de ediciones extraordinarias del {ayer:%d/%m/%Y}")
        normas = extra + normas
    print(f"  {len(normas)} normas publicadas.")

    if not args.sin_texto and normas:
        print("  Descargando texto completo de cada norma…")
        with ThreadPoolExecutor(max_workers=4) as ex:
            for i, (n, txt) in enumerate(zip(normas, ex.map(lambda n: descargar_texto(s, n), normas)), 1):
                n.texto = txt
                if i % max(25, len(normas) // 10) == 0 or i == len(normas):
                    print(f"    {i}/{len(normas)}", flush=True)

    seleccion = [n for n in normas if clasificar(n)]
    orden = list(CATEGORIAS)
    seleccion.sort(key=lambda n: (orden.index(n.categorias[0]), n.fecha, n.entidad))
    print(f"  {len(seleccion)} normas de interés.")

    carpeta = Path(cfg["general"].get("carpeta_salida", BASE_DIR / "resumenes"))
    carpeta.mkdir(parents=True, exist_ok=True)
    nombre = (f"Normas_ElPeruano_{desde:%Y%m%d}.docx" if desde == hasta
              else f"Normas_ElPeruano_{desde:%Y%m%d}_{hasta:%Y%m%d}.docx")
    ruta = carpeta / nombre
    generar_word(seleccion, len(normas), desde, hasta, ruta)
    print(f"  Word generado: {ruta}")

    if args.publicar:
        publicar_historial(normas, seleccion, ruta, hasta)

    if args.no_enviar:
        return
    titulo = f"Normas El Peruano – {len(seleccion)} de interés ({desde:%d/%m/%Y}" + \
             ("" if desde == hasta else f" al {hasta:%d/%m/%Y}") + ")"
    try:
        enviar_telegram(cfg, ruta, texto_breve(seleccion, len(normas), desde, hasta), titulo)
    except Exception as e:  # noqa: BLE001
        print(f"[telegram] Error: {e}")


if __name__ == "__main__":
    sys.exit(main())
